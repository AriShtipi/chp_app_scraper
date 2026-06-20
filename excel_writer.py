"""
כתיבת קובץ התוצאות (xlsx) - גיליון השוואת מחירים + גיליון סיכום חיסכון.
"""

import re

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import VAT, TOP_N_LOCAL, TOP_N_ONLINE


def clean_text(s):
    """מסיר תווי Unicode אפס-רוחב ואותיות ASCII שCHP מזריק כהגנת anti-scraping."""
    if not s:
        return s
    s = re.sub(r'[\u200b\u200c\u200d\u200e\u200f\u00ad\u2060\ufeff\u034f\u180e]', '', s)
    s = re.sub(r'[a-zA-Z0-9]', '', s)
    s = re.sub(r' {2,}', ' ', s)
    return s.strip()


def write_results(products, cities, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "תוצאות CHP"
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    BLUE         = "1a73e8"
    WHITE        = "FFFFFF"
    CITY_COLORS  = ["DCE8FB", "D4F0E8", "FEF3D0", "F3E8FF", "FFE8E8"]
    LOCAL_HDR    = "4A90D9"
    ONLINE_HDR   = "27AE60"
    CAT_BG       = "E8F0FE"

    def c(r, col, val="", bold=False, bg=None, fg="000000",
          align="right", fmt=None, wrap=False, size=10):
        cl = ws.cell(row=r, column=col, value=val)
        cl.font = Font(name="Arial", bold=bold, color=fg, size=size)
        if bg:
            cl.fill = PatternFill("solid", fgColor=bg)
        cl.alignment = Alignment(horizontal=align, vertical="center",
                                  wrap_text=wrap, readingOrder=2)
        cl.border = border
        if fmt:
            cl.number_format = fmt
        return cl

    FIXED = 5
    cols_per_city = TOP_N_LOCAL * 2 + 1 + TOP_N_ONLINE * 2 + 1
    city_starts = {}
    col_idx = FIXED + 1
    for city in cities:
        city_starts[city["name"]] = col_idx
        col_idx += cols_per_city
    TOTAL_COLS = col_idx - 1

    # ── שורה 1: כותרת ──
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COLS)
    cl = ws.cell(row=1, column=1, value="📊  השוואת מחירי CHP — תוצאות סריקה")
    cl.font = Font(name="Arial", bold=True, size=13, color=WHITE)
    cl.fill = PatternFill("solid", fgColor=BLUE)
    cl.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
    ws.row_dimensions[1].height = 26

    # ── שורה 2: headers קבועים + כותרות עיר ──
    fixed_headers = [
        "קטגוריה", "שם פריט", "ברקוד",
        'מחיר קנייה\n(ללא מע"מ)',
        f'מחיר קנייה\n(+מע"מ {round((VAT-1)*100)}%)'
    ]
    for ci, h in enumerate(fixed_headers, 1):
        c(2, ci, h, bold=True, bg=BLUE, fg=WHITE, align="center", wrap=True)

    for i, city in enumerate(cities):
        start = city_starts[city["name"]]
        ws.merge_cells(start_row=2, start_column=start,
                        end_row=2, end_column=start + cols_per_city - 1)
        cl = ws.cell(row=2, column=start, value=city["name"])
        cl.font = Font(name="Arial", bold=True, size=11, color="222222")
        cl.fill = PatternFill("solid", fgColor=CITY_COLORS[i % len(CITY_COLORS)])
        cl.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
        cl.border = border
    ws.row_dimensions[2].height = 22

    # ── שורה 3: sub-headers (מקומי | אונליין) ──
    for ci in range(1, FIXED + 1):
        c(3, ci, "", bg="F0F4FF")

    for i, city in enumerate(cities):
        col_idx = city_starts[city["name"]]
        local_span = TOP_N_LOCAL * 2 + 1
        ws.merge_cells(start_row=3, start_column=col_idx,
                        end_row=3, end_column=col_idx + local_span - 1)
        cl = ws.cell(row=3, column=col_idx, value="🏪 חנויות פיזיות")
        cl.font = Font(name="Arial", bold=True, size=10, color=WHITE)
        cl.fill = PatternFill("solid", fgColor=LOCAL_HDR)
        cl.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
        cl.border = border
        col_idx += local_span

        online_span = TOP_N_ONLINE * 2 + 1
        ws.merge_cells(start_row=3, start_column=col_idx,
                        end_row=3, end_column=col_idx + online_span - 1)
        cl = ws.cell(row=3, column=col_idx, value="🌐 אונליין")
        cl.font = Font(name="Arial", bold=True, size=10, color=WHITE)
        cl.fill = PatternFill("solid", fgColor=ONLINE_HDR)
        cl.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
        cl.border = border
        col_idx += online_span
    ws.row_dimensions[3].height = 20

    # ── שורה 4: field sub-headers ──
    for ci in range(1, FIXED + 1):
        c(4, ci, "", bg="F0F4FF")

    for i, city in enumerate(cities):
        col_idx = city_starts[city["name"]]
        for n in range(TOP_N_LOCAL):
            c(4, col_idx,   f"#{n+1} שם עסק", bold=True, bg="EBF3FB", align="center")
            c(4, col_idx+1, f"#{n+1} מחיר",   bold=True, bg="EBF3FB", align="center")
            col_idx += 2
        c(4, col_idx, "הכי זול", bold=True, bg="EBF3FB", align="center")
        col_idx += 1
        for n in range(TOP_N_ONLINE):
            c(4, col_idx,   f"#{n+1} אתר",    bold=True, bg="E8F8F1", align="center")
            c(4, col_idx+1, f"#{n+1} מחיר",   bold=True, bg="E8F8F1", align="center")
            col_idx += 2
        c(4, col_idx, "הכי זול", bold=True, bg="E8F8F1", align="center")
        col_idx += 1
    ws.row_dimensions[4].height = 20

    # ── שורות נתונים ──
    row_num = 5
    last_cat = ""

    def color_min(cl_cell, effective, inc_vat):
        if not effective or not inc_vat:
            return
        pct = (inc_vat - effective) / inc_vat * 100
        if pct >= 25:
            cl_cell.fill = PatternFill("solid", fgColor="D4EDDA")
            cl_cell.font = Font(name="Arial", bold=True, color="155724", size=10)
        elif pct >= 10:
            cl_cell.fill = PatternFill("solid", fgColor="FFF3CD")
            cl_cell.font = Font(name="Arial", bold=True, color="856404", size=10)
        else:
            cl_cell.fill = PatternFill("solid", fgColor="F8D7DA")
            cl_cell.font = Font(name="Arial", bold=True, color="721c24", size=10)

    def get_global_min(prod):
        """מחזיר את המחיר הכי זול מכל הערים והסוגים."""
        mins = []
        for city in cities:
            city_data = prod["chp"].get(city["name"], {})
            for entry in city_data.get("local", [])[:1]:
                mins.append(entry["effective"])
            for entry in city_data.get("online", [])[:1]:
                mins.append(entry["effective"])
        return min(mins) if mins else None

    for prod in products:
        # שורת קטגוריה
        if prod["category"] != last_cat and prod["category"]:
            ws.merge_cells(start_row=row_num, start_column=1,
                            end_row=row_num, end_column=TOTAL_COLS)
            cl = ws.cell(row=row_num, column=1, value=f"  {prod['category']}")
            cl.font = Font(name="Arial", bold=True, size=10, color="1a3a6b")
            cl.fill = PatternFill("solid", fgColor=CAT_BG)
            cl.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
            cl.border = border
            ws.row_dimensions[row_num].height = 18
            row_num += 1
            last_cat = prod["category"]

        c(row_num, 1, prod["category"])
        c(row_num, 2, prod["name"])
        c(row_num, 3, prod["barcode"])
        c(row_num, 4, prod["price_buy_ex_vat"],  fmt='0.00')
        c(row_num, 5, prod["price_buy_inc_vat"],  fmt='0.00', bold=True)

        inc_vat    = prod["price_buy_inc_vat"]
        global_min = get_global_min(prod)
        col_idx    = FIXED + 1

        for city in cities:
            city_data     = prod["chp"].get(city["name"], {"local": [], "online": []})
            local_prices  = city_data.get("local",  [])[:TOP_N_LOCAL]
            online_prices = city_data.get("online", [])[:TOP_N_ONLINE]

            # עמודות מקומי
            local_min = None
            for n in range(TOP_N_LOCAL):
                if n < len(local_prices):
                    e = local_prices[n]
                    network_clean = clean_text(e.get("network") or "")
                    store_clean   = clean_text(e.get("store")   or "")
                    if network_clean and store_clean and network_clean != store_clean:
                        label = f"{network_clean} - {store_clean}"
                    else:
                        label = network_clean or store_clean or "—"
                    pval = e["effective"]
                    if local_min is None or pval < local_min:
                        local_min = pval

                    is_global_best = global_min and abs(pval - global_min) < 0.001
                    bg_store = "C8F7C5" if is_global_best else "FAFAFA"
                    bg_price = "C8F7C5" if is_global_best else "FAFAFA"
                    c(row_num, col_idx,   label, bg=bg_store)
                    c(row_num, col_idx+1, pval,  bg=bg_price, fmt='0.00')
                else:
                    c(row_num, col_idx,   "—", bg="FAFAFA")
                    c(row_num, col_idx+1, "—", bg="FAFAFA")
                col_idx += 2

            min_cell = c(row_num, col_idx,
                         local_min if local_min else "לא נמצא",
                         fmt='0.00' if local_min else None, bold=bool(local_min))
            if local_min:
                color_min(min_cell, local_min, inc_vat)
            col_idx += 1

            # עמודות אונליין
            online_min = None
            for n in range(TOP_N_ONLINE):
                if n < len(online_prices):
                    e = online_prices[n]
                    network_clean = clean_text(e.get("network") or "")
                    store_clean   = clean_text(e.get("store")   or "")
                    if network_clean and store_clean and network_clean != store_clean:
                        label = f"{network_clean} - {store_clean}"
                    else:
                        label = network_clean or store_clean or "—"
                    pval = e["effective"]
                    if online_min is None or pval < online_min:
                        online_min = pval

                    is_global_best = global_min and abs(pval - global_min) < 0.001
                    bg_store = "C8F7C5" if is_global_best else "F0FFF8"
                    bg_price = "C8F7C5" if is_global_best else "F0FFF8"
                    c(row_num, col_idx,   label, bg=bg_store)
                    c(row_num, col_idx+1, pval,  bg=bg_price, fmt='0.00')
                else:
                    c(row_num, col_idx,   "—", bg="F0FFF8")
                    c(row_num, col_idx+1, "—", bg="F0FFF8")
                col_idx += 2

            min_cell = c(row_num, col_idx,
                         online_min if online_min else "לא נמצא",
                         fmt='0.00' if online_min else None, bold=bool(online_min))
            if online_min:
                color_min(min_cell, online_min, inc_vat)
            col_idx += 1

        ws.row_dimensions[row_num].height = 18
        row_num += 1

    _write_summary_sheet(wb, products, cities)

    # רוחב עמודות
    widths = [14, 36, 16, 12, 12]
    for _ in cities:
        for _ in range(TOP_N_LOCAL):
            widths += [22, 9]
        widths.append(10)
        for _ in range(TOP_N_ONLINE):
            widths += [22, 9]
        widths.append(10)
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = "A5"
    wb.save(output_path)


def _write_summary_sheet(wb, products, cities):
    """
    גיליון סיכום נפרד — לכל מוצר: מחיר קנייה, המחיר הכי זול שנמצא,
    חיסכון %, ועיר/סוג מקור.
    """
    ws = wb.create_sheet("סיכום חיסכון")
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def c(r, col, val="", bold=False, bg=None, fg="000000", align="right", fmt=None):
        cl = ws.cell(row=r, column=col, value=val)
        cl.font = Font(name="Arial", bold=bold, color=fg, size=10)
        if bg:
            cl.fill = PatternFill("solid", fgColor=bg)
        cl.alignment = Alignment(horizontal=align, vertical="center", readingOrder=2)
        cl.border = border
        if fmt:
            cl.number_format = fmt
        return cl

    headers = ["שם פריט", "ברקוד", 'מחיר קנייה (+מע"מ)', "מחיר כי זול שנמצא", "חיסכון ₪", "חיסכון %", "מקור"]
    BLUE, WHITE = "1a73e8", "FFFFFF"
    for ci, h in enumerate(headers, 1):
        c(1, ci, h, bold=True, bg=BLUE, fg=WHITE, align="center")
    ws.row_dimensions[1].height = 22

    col_widths = [36, 16, 18, 18, 12, 12, 30]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    row_num = 2
    for prod in products:
        inc_vat    = prod["price_buy_inc_vat"]
        best_price = None
        best_src   = "—"

        for city in cities:
            city_data = prod["chp"].get(city["name"], {})
            for e in city_data.get("local", [])[:1]:
                if best_price is None or e["effective"] < best_price:
                    best_price = e["effective"]
                    store = clean_text(e.get("network") or e.get("store") or "")
                    best_src = f"{city['name'].strip()} / {store}"
            for e in city_data.get("online", [])[:1]:
                if best_price is None or e["effective"] < best_price:
                    best_price = e["effective"]
                    store = clean_text(e.get("network") or e.get("store") or "")
                    best_src = f"אונליין / {store}"

        saving_nis = round(inc_vat - best_price, 2) if best_price else None
        saving_pct = round((inc_vat - best_price) / inc_vat * 100, 1) if best_price and inc_vat else None

        c(row_num, 1, prod["name"])
        c(row_num, 2, prod["barcode"])
        c(row_num, 3, inc_vat,    fmt='0.00', bold=True)
        c(row_num, 4, best_price if best_price else "לא נמצא", fmt='0.00' if best_price else None)
        c(row_num, 5, saving_nis if saving_nis is not None else "—", fmt='0.00' if saving_nis else None)

        pct_cell = c(row_num, 6, saving_pct if saving_pct is not None else "—",
                     fmt='0.0' if saving_pct else None)
        if saving_pct is not None:
            if saving_pct >= 25:
                pct_cell.fill = PatternFill("solid", fgColor="D4EDDA")
                pct_cell.font = Font(name="Arial", bold=True, color="155724", size=10)
            elif saving_pct >= 10:
                pct_cell.fill = PatternFill("solid", fgColor="FFF3CD")
                pct_cell.font = Font(name="Arial", bold=True, color="856404", size=10)
            else:
                pct_cell.fill = PatternFill("solid", fgColor="F8D7DA")
                pct_cell.font = Font(name="Arial", bold=True, color="721c24", size=10)

        c(row_num, 7, best_src)
        ws.row_dimensions[row_num].height = 17
        row_num += 1

    ws.freeze_panes = "A2"