"""
קריאת קובצי ספק (xlsx) והמרתם לרשימת מוצרים.
תומך בשני פורמטים אוטומטית:
  1. פורמט שטראוס/פריטו-לי  — header בשורה 1, ברקוד בעמודה A
  2. פורמט legacy            — header מוסתר, ברקוד בעמודה B (int >= 12 ספרות)
"""

import openpyxl

from config import VAT


def _is_barcode(val) -> bool:
    """מחזיר True אם הערך נראה כמו ברקוד (מינימום 8 ספרות)."""
    if val is None:
        return False
    s = str(val).strip().split(".")[0]
    return s.isdigit() and len(s) >= 8


def _detect_format(ws) -> str:
    """
    מזהה אוטומטית את פורמט הקובץ:
      'strauss' — פריטו-לי / שטראוס: header בשורה 1, ברקוד בעמודה A (string)
      'legacy'  — הפורמט הישן: header מוסתר באמצע, ברקוד בעמודה B (int >= 12 ספרות)
    """
    first_cell = str(ws.cell(row=1, column=1).value or "")
    if "ברקוד" in first_cell:
        return "strauss"
    return "legacy"


def _read_strauss_format(ws) -> list:
    """
    פורמט שטראוס / פריטו-לי:
      שורה 1  — כותרות עמודות (ברקוד יחידה | חומר | שם חומר | משקל | כשרות | מחיר | כמות בקרטון)
      שורה 2+ — נתונים
      אין שורות קטגוריה, אין עמודת מחיר קנייה נפרדת
    """
    products = []

    # מיפוי דינמי לפי שמות עמודות
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    col_barcode = next((v for k, v in headers.items() if k and "ברקוד" in str(k)), 1)
    col_name    = next((v for k, v in headers.items() if k and "שם" in str(k)), 3)
    col_price   = next((v for k, v in headers.items() if k and "מחיר" in str(k)), 6)

    for row_idx in range(2, ws.max_row + 1):
        barcode_val = ws.cell(row=row_idx, column=col_barcode).value
        if not _is_barcode(barcode_val):
            continue

        barcode   = str(barcode_val).strip().split(".")[0]
        name      = str(ws.cell(row=row_idx, column=col_name).value or "").strip().strip('"\'')
        price_raw = ws.cell(row=row_idx, column=col_price).value
        price     = round(float(price_raw), 4) if isinstance(price_raw, (int, float)) else 0.0

        products.append({
            "category":          "",
            "barcode":           barcode,
            "name":              name,
            "price_list_ex_vat": price,
            "price_buy_ex_vat":  price,
            "price_buy_inc_vat": round(price * VAT, 2),
            "chp":               {}
        })

    return products


def _read_legacy_format(ws) -> list:
    """
    הפורמט הישן:
      - header מוסתר שמכיל "ברקוד" — מסמן תחילת נתונים
      - עמודה B (index 1) = ברקוד כ-int >= 12 ספרות
      - עמודה C (index 2) = שם מוצר
      - עמודה D (index 3) = מחיר מחירון (ללא מע"מ)
      - עמודה F (index 5) = מחיר קנייה (ללא מע"מ)
      - שורות טקסט = שורות קטגוריה
    """
    products = []
    current_category = ""
    header_found = False

    for row in ws.iter_rows(values_only=True):
        if not header_found:
            row_text = " ".join(str(v or "") for v in row)
            if "ברקוד" in row_text:
                header_found = True
            elif any(v and isinstance(v, str) and len(v.strip()) > 3
                     and not str(v).startswith("יח") for v in row):
                cat = next((str(v).strip() for v in row
                            if v and isinstance(v, str) and len(str(v).strip()) > 3), "")
                if cat:
                    current_category = cat
            continue

        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        d = row[3] if len(row) > 3 else None
        f = row[5] if len(row) > 5 else None

        if b and isinstance(b, (int, float)) and len(str(int(b))) >= 12:
            barcode    = str(int(b))
            name       = str(c or "").strip().strip('"\'')
            price_list = round(float(d), 4) if isinstance(d, (int, float)) else 0.0
            price_buy  = round(float(f), 4) if isinstance(f, (int, float)) else price_list
            products.append({
                "category":          current_category,
                "barcode":           barcode,
                "name":              name,
                "price_list_ex_vat": price_list,
                "price_buy_ex_vat":  price_buy,
                "price_buy_inc_vat": round(price_buy * VAT, 2),
                "chp":               {}
            })
        else:
            for v in row:
                if v and isinstance(v, str):
                    t = v.strip().strip('"\'')
                    if t and len(t) > 2 and not t.startswith("יח"):
                        current_category = t
                        break

    return products


def read_supplier_file(path: str) -> list:
    """
    קורא קובץ ספק ומחזיר רשימת מוצרים.
    תומך בשני פורמטים אוטומטית:
      1. פורמט שטראוס/פריטו-לי  — header בשורה 1, ברקוד בעמודה A
      2. פורמט legacy            — header מוסתר, ברקוד בעמודה B (int >= 12 ספרות)
    """
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    fmt = _detect_format(ws)
    if fmt == "strauss":
        return _read_strauss_format(ws)
    return _read_legacy_format(ws)